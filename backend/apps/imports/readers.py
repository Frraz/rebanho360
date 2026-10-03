"""Leitor de planilha (F2-09): openpyxl com as regras que vieram da análise.

- Data pode vir como `datetime` ou como serial do Excel (`45840` =
  2025-07-02, origem 1899-12-30).
- **`#DIV/0!`, `#N/A` e célula vazia viram pendência tratável, nunca falha.**
- **Linha em branco não é lançamento de R$ 0,00**: sem data e sem valor, a
  linha não existe, mesmo com outras células preenchidas (pagador arrastado).
- `-` é o marcador de "vazio" que a planilha usa (FRETE, COMISSÃO...).

O que sai daqui é JSON puro (`raw` de `ImportRow`): datas viram `AAAA-MM-DD`,
erros de célula viram `{"erro": "#DIV/0!"}`. Nada é gravado em tabela final.
"""

import datetime
import re
import unicodedata
from dataclasses import dataclass, field
from decimal import Decimal, InvalidOperation

import openpyxl

ORIGEM_DO_EXCEL = datetime.date(1899, 12, 30)
ERROS_DE_CELULA = {"#DIV/0!", "#N/A", "#VALUE!", "#REF!", "#NAME?", "#NUM!", "#NULL!"}
# Serial plausível de data de operação pecuária (1990..2100): evita tomar
# a quantidade "133" por uma data de 1900.
_SERIAL_MIN, _SERIAL_MAX = 32874, 73415


def normalizar(texto) -> str:
    """Sem acento, minúsculo, espaços colapsados — para casar nomes que a
    planilha e o cadastro grafam diferente ("Femeas" × "Fêmeas")."""
    if texto is None:
        return ""
    sem_acento = unicodedata.normalize("NFKD", str(texto))
    sem_acento = "".join(c for c in sem_acento if not unicodedata.combining(c))
    return re.sub(r"\s+", " ", sem_acento).strip().lower()


# --------------------------------------------------------------------------
# Células
# --------------------------------------------------------------------------


def valor_da_celula(valor):
    """Célula do openpyxl → valor JSON-safe. Vazio e `-` → `None`."""
    if valor is None:
        return None
    if isinstance(valor, datetime.datetime):
        return valor.date().isoformat()
    if isinstance(valor, datetime.date):
        return valor.isoformat()
    if isinstance(valor, bool):
        return valor
    if isinstance(valor, int | float):
        return valor
    if isinstance(valor, Decimal):
        return str(valor)
    texto = str(valor).strip()
    if texto in ("", "-", "–"):
        return None
    if texto.upper() in ERROS_DE_CELULA:
        return {"erro": texto.upper()}
    return texto


def eh_erro_de_celula(valor) -> bool:
    return isinstance(valor, dict) and "erro" in valor


def tem_valor(valor) -> bool:
    """Há conteúdo utilizável (nem vazio, nem erro de célula)."""
    return valor is not None and not eh_erro_de_celula(valor)


def texto_da_celula(valor) -> str:
    if not tem_valor(valor):
        return ""
    return str(valor).strip()


def data_da_celula(valor) -> datetime.date | None:
    """ISO, `dd/mm/aaaa` ou serial do Excel. Qualquer outra coisa → `None`."""
    if not tem_valor(valor) or isinstance(valor, bool):
        return None
    if isinstance(valor, int | float):
        if not _SERIAL_MIN <= valor <= _SERIAL_MAX:
            return None
        return ORIGEM_DO_EXCEL + datetime.timedelta(days=int(valor))
    texto = str(valor).strip()
    for formato in ("%Y-%m-%d", "%d/%m/%Y"):
        try:
            return datetime.datetime.strptime(texto, formato).date()
        except ValueError:
            continue
    return None


def decimal_da_celula(valor) -> Decimal | None:
    """Número da planilha → `Decimal`, passando por `str` (nunca `float`
    direto: 0.1 + 0.2 não é 0.3). Texto aceita "1.234,56" e "1234.56"."""
    if not tem_valor(valor) or isinstance(valor, bool):
        return None
    try:
        if isinstance(valor, int | float):
            return Decimal(str(valor))
        texto = str(valor).strip().replace("R$", "").replace(" ", "")
        if "," in texto:
            texto = texto.replace(".", "").replace(",", ".")
        return Decimal(texto)
    except InvalidOperation:
        return None


def inteiro_da_celula(valor) -> int | None:
    numero = decimal_da_celula(valor)
    if numero is None or numero != numero.to_integral_value():
        return None
    return int(numero)


# --------------------------------------------------------------------------
# Abertura
# --------------------------------------------------------------------------


class PlanilhaInvalida(Exception):
    """Arquivo que não é uma planilha legível, ou sem a aba esperada."""


def abrir_planilha(arquivo):
    """`data_only=True`: lê o valor calculado, não a fórmula."""
    try:
        return openpyxl.load_workbook(arquivo, data_only=True, read_only=True)
    except Exception as exc:  # zip corrompido, formato errado...
        raise PlanilhaInvalida(
            "Não consegui ler este arquivo como planilha Excel (.xlsx)."
        ) from exc


def aba(workbook, nome: str):
    for titulo in workbook.sheetnames:
        if normalizar(titulo) == normalizar(nome):
            return workbook[titulo]
    raise PlanilhaInvalida(
        f"A planilha não tem a aba {nome}. Abas encontradas: "
        + ", ".join(workbook.sheetnames)
    )


@dataclass
class LinhaLida:
    sheet: str
    row_number: int
    raw: dict


@dataclass
class Leitura:
    linhas: list[LinhaLida] = field(default_factory=list)
    stats: dict = field(default_factory=dict)


def _linhas(ws, *, desde: int = 1):
    """(número da linha, valores da linha já normalizados)."""
    for numero, linha in enumerate(
        ws.iter_rows(min_row=desde, values_only=True), start=desde
    ):
        yield numero, [valor_da_celula(c) for c in linha]


def _achar_cabecalho(ws, *, obrigatorias: set[str], limite: int = 30):
    """Primeira linha (até `limite`) que tem todas as colunas esperadas.
    Devolve `(número, {nome normalizado: índice})`."""
    for numero, valores in _linhas(ws):
        if numero > limite:
            break
        nomes: dict[str, int] = {}
        for indice, valor in enumerate(valores):
            if isinstance(valor, str):
                # A primeira ocorrência vale: a aba COMPRA DE GADO tem outras
                # tabelas à direita, com colunas de mesmo nome.
                nomes.setdefault(normalizar(valor), indice)
        if obrigatorias <= set(nomes):
            return numero, nomes
    raise PlanilhaInvalida(
        f"Não encontrei o cabeçalho esperado na aba {ws.title} "
        f"(colunas: {', '.join(sorted(obrigatorias))})."
    )


def _celula(valores, indice):
    return valores[indice] if indice is not None and indice < len(valores) else None


# --------------------------------------------------------------------------
# CUSTOS
# --------------------------------------------------------------------------

COLUNAS_CUSTOS = {
    "data": "data",
    "pagador": "pagador",
    "item": "item",
    "valor": "valor total",
    "sub_centro": "sub centro",
    "centro": "centro de custo",
    "classe": "classe",
    # `MÊS` e `ANO` são digitadas e redundantes com `DATA` — não viram campo.
    # O ANO só é lido para apontar data com o ano digitado errado.
    "ano": "ano",
}


def ler_custos(workbook) -> Leitura:
    """Aba `CUSTOS`: 417 linhas lidas → 235 lançamentos reais."""
    ws = aba(workbook, "CUSTOS")
    cabecalho, colunas = _achar_cabecalho(ws, obrigatorias={"data", "valor total"})
    leitura = Leitura(stats={"lidas": 0, "em_branco": 0, "lancamentos": 0})

    for numero, valores in _linhas(ws, desde=cabecalho + 1):
        leitura.stats["lidas"] += 1
        raw = {
            chave: _celula(valores, colunas.get(nome))
            for chave, nome in COLUNAS_CUSTOS.items()
        }
        # Sem data e sem valor a linha não existe — mesmo com o pagador
        # arrastado para baixo. Nunca vira lançamento de R$ 0,00.
        if not tem_valor(raw["data"]) and not tem_valor(raw["valor"]):
            leitura.stats["em_branco"] += 1
            continue
        leitura.stats["lancamentos"] += 1
        leitura.linhas.append(LinhaLida("CUSTOS", numero, raw))
    return leitura


# --------------------------------------------------------------------------
# COMPRA DE GADO
# --------------------------------------------------------------------------

COLUNAS_COMPRAS = {
    "data": "data",
    "quantidade": "quant.",
    "categoria": "novilhas / bois",
    "valor": "valor total da compra",
    "fazenda": "fazenda",
    "parceria": "parceria",
    "frete": "frete",
    "comissao": "comissao",
    "observacoes": "observacoes",
    "impostos": "impostos",
}


def ler_compras(workbook) -> Leitura:
    """Aba `COMPRA DE GADO`: 13 compras. `MÉDIA / CAB` é derivada (com
    `#DIV/0!`) e nem é lida."""
    ws = aba(workbook, "COMPRA DE GADO")
    cabecalho, colunas = _achar_cabecalho(ws, obrigatorias={"data", "quant."})
    leitura = Leitura(stats={"lidas": 0, "em_branco": 0, "compras": 0})

    for numero, valores in _linhas(ws, desde=cabecalho + 1):
        if not any(v is not None for v in valores):
            continue
        leitura.stats["lidas"] += 1
        raw = {
            chave: _celula(valores, colunas.get(nome))
            for chave, nome in COLUNAS_COMPRAS.items()
        }
        # Linha de fórmula sem dado: só `#DIV/0!` na média.
        if not tem_valor(raw["data"]) and not tem_valor(raw["quantidade"]):
            leitura.stats["em_branco"] += 1
            continue
        leitura.stats["compras"] += 1
        leitura.linhas.append(LinhaLida("COMPRA DE GADO", numero, raw))
    return leitura


# --------------------------------------------------------------------------
# Abas de fazenda
# --------------------------------------------------------------------------

#: Abas de fazenda (nome na planilha). `GERAL` é consolidação derivada — e é
#: a que fecha em −140: descartada.
ABAS_DE_FAZENDA = [
    "SÃO FRANCISCO",
    "GOIANO",
    "BAIXÃO",
    "MORADA DO BOI",
    "SAO JOSE DO GROTAO",
]

COLUNAS_MOVIMENTO = {
    "data": "data",
    "categoria": "categoria",
    "tipo": "tipo",
    "quantidade": "quantidade",
    "saida": "saida",
    "destino": "destino",
    "peso_total": "peso total",
}


def ler_aba_de_fazenda(workbook, nome: str) -> Leitura:
    """Uma aba de fazenda tem duas partes: o quadro-resumo (linhas 4-16),
    **derivado**, e o lançamento real (linhas 19+).

    O quadro-resumo é descartado — com uma exceção que a spec não previu: a
    coluna `SALDO ANTERIOR` é dado de **entrada**, não derivado. Sem ela as
    mortes e os abates deixariam o saldo negativo. Vira linha `SALDO ANTERIOR`.
    """
    ws = aba(workbook, nome)
    leitura = Leitura(
        stats={"lidas": 0, "em_branco": 0, "movimentos": 0, "saldos_anteriores": 0}
    )

    todas = list(_linhas(ws))
    # --- quadro-resumo: só o SALDO ANTERIOR ---
    cab_resumo = next(
        (
            (n, v)
            for n, v in todas
            if normalizar(_celula(v, 0)) == "categorias"
            and any(normalizar(x) == "saldo anterior" for x in v if isinstance(x, str))
        ),
        None,
    )
    if cab_resumo:
        numero_cab, valores_cab = cab_resumo
        coluna_saldo = next(
            i for i, x in enumerate(valores_cab) if normalizar(x) == "saldo anterior"
        )
        for numero, valores in todas:
            if numero <= numero_cab:
                continue
            rotulo = texto_da_celula(_celula(valores, 0))
            if normalizar(rotulo) == "total":
                break
            saldo = _celula(valores, coluna_saldo)
            if inteiro_da_celula(saldo):
                leitura.stats["saldos_anteriores"] += 1
                leitura.linhas.append(
                    LinhaLida(
                        ws.title,
                        numero,
                        {
                            "tipo": "SALDO ANTERIOR",
                            "categoria": rotulo,
                            "quantidade": saldo,
                            "origem": "quadro-resumo da aba (SALDO ANTERIOR)",
                        },
                    )
                )

    # --- lançamento real ---
    cabecalho, colunas = _achar_cabecalho(
        ws, obrigatorias={"data", "categoria", "tipo", "quantidade"}
    )
    for numero, valores in todas:
        if numero <= cabecalho:
            continue
        leitura.stats["lidas"] += 1
        raw = {
            chave: _celula(valores, colunas.get(nome_coluna))
            for chave, nome_coluna in COLUNAS_MOVIMENTO.items()
        }
        if not tem_valor(raw["data"]) and not tem_valor(raw["quantidade"]):
            leitura.stats["em_branco"] += 1
            continue
        leitura.stats["movimentos"] += 1
        leitura.linhas.append(LinhaLida(ws.title, numero, raw))
    return leitura


def ler_movimentacoes(workbook) -> Leitura:
    """Todas as abas de fazenda que existirem na planilha."""
    resultado = Leitura(stats={"abas": {}})
    nomes = {normalizar(t): t for t in workbook.sheetnames}
    for nome in ABAS_DE_FAZENDA:
        titulo = nomes.get(normalizar(nome))
        if titulo is None:
            continue
        leitura = ler_aba_de_fazenda(workbook, titulo)
        resultado.linhas.extend(leitura.linhas)
        resultado.stats["abas"][titulo] = leitura.stats
    if not resultado.stats["abas"]:
        raise PlanilhaInvalida(
            "Nenhuma aba de fazenda encontrada. Esperava: " + ", ".join(ABAS_DE_FAZENDA)
        )
    return resultado


# --------------------------------------------------------------------------
# VENDAS
# --------------------------------------------------------------------------

COLUNAS_VENDAS = {
    "data": "data",
    "tipo": "tipo de venda",
    "categoria": "categoria",
    "animais": "animais",
    "peso_total": "peso total",
    "carcaca_total": "carcaca total",
    "valor_total": "valor total",
    "comprador": "comprador",
    "forma": "forma",
    "fazenda": "fazenda",
    "parceria": "parceria",
    "safra": "safra",
}

#: Os derivados que a planilha grava como valor. **Não viram campo**: são
#: lidos só para a conferência — o `CarcassService` recalcula e a diferença
#: vira aviso (foi assim que a pendência #7 apareceu).
COLUNAS_DERIVADAS_DE_VENDAS = {
    "peso_medio": "peso medio",
    "carcaca_media": "carcaca media",
    "rendimento": "rendimento %",
    "valor_cabeca": "valor cabeca",
    "valor_arroba": "valor por @",
    "soma_rendimento": "soma rendimento",
}


def ler_vendas(workbook) -> Leitura:
    """Aba `VENDAS`: 3 abates, 354 cabeças. As linhas abaixo das vendas só
    têm `#DIV/0!` nos derivados — fórmulas sem dado, não vendas."""
    ws = aba(workbook, "VENDAS")
    cabecalho, colunas = _achar_cabecalho(
        ws, obrigatorias={"data", "tipo de venda", "animais", "valor total"}
    )
    leitura = Leitura(stats={"lidas": 0, "em_branco": 0, "vendas": 0})

    for numero, valores in _linhas(ws, desde=cabecalho + 1):
        if not any(v is not None for v in valores):
            continue
        leitura.stats["lidas"] += 1
        raw = {
            chave: _celula(valores, colunas.get(nome))
            for chave, nome in COLUNAS_VENDAS.items()
        }
        if not tem_valor(raw["data"]) and not tem_valor(raw["animais"]):
            leitura.stats["em_branco"] += 1
            continue
        raw["planilha"] = {
            chave: _celula(valores, colunas.get(nome))
            for chave, nome in COLUNAS_DERIVADAS_DE_VENDAS.items()
        }
        leitura.stats["vendas"] += 1
        leitura.linhas.append(LinhaLida("VENDAS", numero, raw))
    return leitura


# --------------------------------------------------------------------------
# PESAGENS E CONFERENCIA — o desempilhamento
# --------------------------------------------------------------------------

ABA_DE_PESAGENS = "PESAGENS E CONFERENCIA"


def _blocos_de_pesagem(valores) -> list[int]:
    """Índice da coluna `DATA` de cada bloco `(DATA, BRINCO, PESO,
    MOVIMENTAÇÃO)` da linha de cabeçalho. Achados pelo cabeçalho, não por
    letras fixas: a doc dizia 8 blocos, a planilha real tem 9."""
    nomes = [normalizar(v) if isinstance(v, str) else "" for v in valores]
    return [
        i
        for i in range(len(nomes) - 3)
        if nomes[i] == "data"
        and nomes[i + 1] == "brinco"
        and nomes[i + 2] == "peso"
        and nomes[i + 3].startswith("movimenta")
    ]


def ler_pesagens(workbook) -> Leitura:
    """Oito (nove) blocos paralelos viram uma tabela só: cada animal
    pesado é uma linha. **Não existe entidade `Animal`** (ADR 0004): o
    brinco é só um texto da pesagem."""
    ws = aba(workbook, ABA_DE_PESAGENS)
    todas = list(_linhas(ws))
    cabecalho, blocos = None, []
    for numero, valores in todas:
        if numero > 15:
            break
        blocos = _blocos_de_pesagem(valores)
        if blocos:
            cabecalho = numero
            break
    if cabecalho is None:
        raise PlanilhaInvalida(
            f"Não encontrei os blocos (DATA, BRINCO, PESO, MOVIMENTAÇÃO) na aba {ws.title}."
        )

    leitura = Leitura(
        stats={"lidas": 0, "em_branco": 0, "blocos": len(blocos), "animais": 0}
    )
    for numero, valores in todas:
        if numero <= cabecalho:
            continue
        for inicio in blocos:
            data, brinco, peso, movimento = (
                _celula(valores, inicio + i) for i in range(4)
            )
            if not any(tem_valor(x) for x in (data, brinco, peso)):
                continue
            letra = openpyxl.utils.get_column_letter(inicio + 1)
            leitura.stats["lidas"] += 1
            leitura.stats["animais"] += 1
            leitura.linhas.append(
                LinhaLida(
                    f"{ABA_DE_PESAGENS} · bloco {letra}",
                    numero,
                    {
                        "data": data,
                        "brinco": brinco,
                        "peso": peso,
                        "movimentacao": movimento,
                        "bloco": letra,
                    },
                )
            )
    return leitura
