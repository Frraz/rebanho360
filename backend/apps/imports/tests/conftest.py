import datetime
import io
from pathlib import Path

import openpyxl
import pytest
from django.core.files.uploadedfile import SimpleUploadedFile
from django.core.management import call_command

from apps.accounts.models import User
from apps.costs.tests.conftest import *  # noqa: F401,F403

RAIZ = Path(__file__).resolve().parents[4]
PLANILHA_REAL = next(
    iter((RAIZ / "docs" / "fontes" / "planilhas").glob("CONTROLE PASTO*.xlsx")), None
)

precisa_da_planilha = pytest.mark.skipif(
    PLANILHA_REAL is None,
    reason="docs/fontes/planilhas/CONTROLE PASTO*.xlsx não está no repositório.",
)


@pytest.fixture(autouse=True)
def midia_temporaria(settings, tmp_path):
    """Arquivos importados vão para uma pasta descartável."""
    settings.MEDIA_ROOT = tmp_path / "media"


@pytest.fixture
def demo(db):
    """Empresa, safra 2025/2026, as 6 fazendas, 11 categorias e um usuário
    por papel — o mesmo que o `seed_demo` entrega ao desenvolvedor."""
    call_command("seed_demo")


@pytest.fixture(autouse=True)
def safra_padrao(demo):
    """Aqui a safra vem do `seed_demo`; sobrescreve a dos testes de custos,
    para não haver duas safras 2025/2026 de empresas diferentes."""
    from apps.organizations.models import Season

    return Season.objects.get(name="2025/2026")


@pytest.fixture
def usuario_escritorio(demo):
    return User.objects.get(username="escritorio@teste")


@pytest.fixture
def usuario_gestor(demo):
    return User.objects.get(username="gestor@teste")


@pytest.fixture
def usuario_campo(demo):
    return User.objects.get(username="campo@teste")


def upload_real() -> SimpleUploadedFile:
    return SimpleUploadedFile(
        PLANILHA_REAL.name,
        PLANILHA_REAL.read_bytes(),
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )


def upload_de(workbook: openpyxl.Workbook, nome="planilha.xlsx") -> SimpleUploadedFile:
    # O .xlsx embute created/modified com resolução de segundos: sem fixar,
    # dois arquivos "iguais" gerados em segundos diferentes têm hashes
    # diferentes — e o teste de reimportação passava só se a máquina fosse
    # rápida (flaky sob carga).
    workbook.properties.created = workbook.properties.modified = datetime.datetime(
        2026, 1, 1
    )
    buffer = io.BytesIO()
    workbook.save(buffer)
    return SimpleUploadedFile(nome, buffer.getvalue())


def aba_de_fazenda(workbook, titulo, *, saldos=(), movimentos=()):
    """Aba no formato das 5 de fazenda: quadro-resumo no topo (cabeçalho na
    linha 4) e lançamento real a partir da linha 19."""
    ws = workbook.create_sheet(titulo)
    ws["A1"] = f"MOVIMENTAÇÃO FAZENDA {titulo}"
    ws["A4"], ws["D4"] = "CATEGORIAS", "SALDO ANTERIOR"
    linha = 5
    for categoria, saldo in saldos:
        ws.cell(linha, 1, categoria)
        ws.cell(linha, 4, saldo)
        linha += 1
    ws.cell(linha, 1, "TOTAL")
    for coluna, nome in enumerate(
        [
            "DATA",
            "MÊS",
            "CATEGORIA",
            "TIPO",
            "QUANTIDADE",
            "SAIDA",
            "DESTINO",
            "PESO TOTAL",
            "PESO MÉDIO",
        ],
        start=1,
    ):
        ws.cell(19, coluna, nome)
    for i, (data, categoria, tipo, quantidade, *resto) in enumerate(
        movimentos, start=20
    ):
        destino = resto[0] if resto else None
        peso = resto[1] if len(resto) > 1 else None
        ws.cell(i, 1, data)
        ws.cell(i, 3, categoria)
        ws.cell(i, 4, tipo)
        ws.cell(i, 5, quantidade)
        ws.cell(i, 7, destino)
        ws.cell(i, 8, peso)
    return ws


def planilha_sem_abas() -> openpyxl.Workbook:
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    return wb
