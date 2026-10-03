"""F3-11 — exportação CSV e XLSX: qualquer relatório exporta com os filtros
aplicados preservados no arquivo. Hoje tudo é planilha: tirar essa saída é
tirar a autonomia de quem analisa por fora."""

import datetime
import io
from decimal import Decimal

import openpyxl
import pytest
from django.urls import reverse

from apps.audit.models import AuditEvent
from apps.reports import services

pytestmark = pytest.mark.django_db

D = Decimal

SLUGS = list(services.RELATORIOS)
# Contas, mapa e pagamentos (Fase 4) não são "da safra do topo": a dívida de uma
# safra se paga em outra. O filtro que imprimem é o que de fato aplicam.
SEM_FILTRO_DE_SAFRA = {
    "contas-a-pagar",
    "contas-a-receber",
    "programacao-de-pagamentos",
    "pagamentos-realizados",
    "mapa-financeiro",
}


def _filtro_esperado(slug):
    return "Filtros aplicados: " + (
        "" if slug in SEM_FILTRO_DE_SAFRA else "Safra 2025/2026"
    )


def abrir(resposta):
    return openpyxl.load_workbook(io.BytesIO(resposta.content))


@pytest.fixture
def com_venda(criar, confirmar, dados_abate):
    return confirmar(criar(**dados_abate))


class TestCsv:
    @pytest.mark.parametrize("slug", SLUGS)
    def test_todo_relatorio_exporta_csv(self, client, gestor, season, slug):
        client.force_login(gestor)
        resposta = client.get(
            reverse("reports:relatorio", args=[slug]) + "?formato=csv"
        )

        assert resposta.status_code == 200
        assert resposta["Content-Type"].startswith("text/csv")
        texto = resposta.content.decode()
        assert _filtro_esperado(slug) in texto and "Todas as fazendas" in texto

    def test_filtros_de_periodo_ficam_dentro_do_csv(
        self, client, gestor, season, com_venda
    ):
        client.force_login(gestor)
        corpo = client.get(
            reverse("reports:relatorio", args=["vendas-e-abates"])
            + "?formato=csv&de=2025-08-01&ate=2025-08-31"
        ).content.decode()

        assert "Período: 01/08/2025 a 31/08/2025" in corpo
        assert "VD-2025/26-0001" in corpo

    def test_notas_vao_no_arquivo(self, client, gestor, season, com_venda):
        client.force_login(gestor)
        corpo = client.get(
            reverse("reports:relatorio", args=["vendas-e-abates"]) + "?formato=csv"
        ).content.decode()
        assert "consideram só as vendas que têm peso de carcaça" in corpo


class TestXlsx:
    @pytest.mark.parametrize("slug", SLUGS)
    def test_todo_relatorio_exporta_xlsx(self, client, gestor, season, slug):
        client.force_login(gestor)
        resposta = client.get(
            reverse("reports:relatorio", args=[slug]) + "?formato=xlsx"
        )

        assert resposta.status_code == 200
        assert resposta["Content-Type"] == (
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        )
        assert resposta["Content-Disposition"].endswith('.xlsx"')
        wb = abrir(resposta)
        ws = wb.active
        assert ws["A1"].value  # título
        textos = [str(c.value) for c in ws["A"] if c.value]
        assert any(t.startswith(_filtro_esperado(slug)) for t in textos)
        assert any(t.startswith("Emitido em") and "por gestor" in t for t in textos)

    def test_os_numeros_sao_numeros_com_formato_nao_texto(
        self, client, gestor, season, com_venda
    ):
        client.force_login(gestor)
        ws = abrir(
            client.get(
                reverse("reports:relatorio", args=["vendas-e-abates"]) + "?formato=xlsx"
            )
        ).active
        cabecalho = next(lin for lin in ws.iter_rows() if lin[0].value == "Venda")
        colunas = {c.value: c.column for c in cabecalho}
        linha = cabecalho[0].row + 1

        valor = ws.cell(linha, colunas["Valor total"])
        assert valor.value == 401502.68 and isinstance(valor.value, int | float)
        assert valor.number_format == '"R$" #,##0.00'
        cabecas = ws.cell(linha, colunas["Cabeças"])
        assert cabecas.value == 84 and cabecas.number_format == "#,##0"
        rendimento = ws.cell(linha, colunas["Rendimento"])
        assert (
            round(rendimento.value, 2) == 51.33
            and rendimento.number_format == '0.00"%"'
        )

    def test_dado_ausente_vai_como_travessao_nao_celula_vazia(
        self, client, gestor, season, criar, confirmar, dados_abate
    ):
        confirmar(criar(**{**dados_abate, "carcass_weight_kg": None}))
        client.force_login(gestor)
        ws = abrir(
            client.get(
                reverse("reports:relatorio", args=["vendas-e-abates"]) + "?formato=xlsx"
            )
        ).active
        cabecalho = next(lin for lin in ws.iter_rows() if lin[0].value == "Venda")
        colunas = {c.value: c.column for c in cabecalho}

        assert ws.cell(cabecalho[0].row + 1, colunas["Rendimento"]).value == "—"

    def test_secoes_viram_abas(self, client, gestor, season, com_venda):
        client.force_login(gestor)
        wb = abrir(
            client.get(
                reverse("reports:relatorio", args=["vendas-e-abates"]) + "?formato=xlsx"
            )
        )
        assert wb.sheetnames == [
            "Vendas e abates",
            "Vendas por mês",
            "Vendas por comprador",
        ]

    def test_filtros_de_periodo_e_lote_ficam_no_arquivo(
        self, client, gestor, season, lote_gordo, escritorio, sao_francisco
    ):
        from apps.herd import services as herd
        from apps.herd.models import WeighingReason

        herd.registrar_pesagem(
            date=datetime.date(2025, 8, 5),
            farm=sao_francisco,
            lot=lote_gordo,
            reason=WeighingReason.CONFERENCIA,
            head_count=10,
            total_weight_kg=D("2000"),
            usuario=escritorio,
        )
        client.force_login(gestor)
        ws = abrir(
            client.get(
                reverse("reports:relatorio", args=["pesagens"])
                + f"?formato=xlsx&de=2025-08-01&ate=2025-08-31&lote={lote_gordo.pk}"
            )
        ).active

        filtros = next(
            str(c.value)
            for c in ws["A"]
            if str(c.value).startswith("Filtros aplicados")
        )
        assert "Período: 01/08/2025 a 31/08/2025" in filtros
        assert "Lote LT-SFR-010" in filtros

    def test_audita_a_exportacao_com_os_filtros(self, client, gestor, season):
        client.force_login(gestor)
        client.get(reverse("reports:relatorio", args=["pesagens"]) + "?formato=xlsx")

        evento = AuditEvent.objects.filter(action="EXPORT", entity_id="pesagens").get()
        assert "xlsx" in evento.reason and "Safra 2025/2026" in evento.reason

    def test_aba_com_titulo_longo_ou_com_caracter_proibido_nao_quebra(self):
        rel = services.Relatorio(
            titulo="Relatório: custos/lote [teste] com um título muito, mas muito comprido",
            descricao="",
            colunas=[services.Coluna("a", "A")],
            linhas=[{"a": "x"}],
        )
        wb = openpyxl.load_workbook(io.BytesIO(services.xlsx_do_relatorio(rel)))
        assert len(wb.sheetnames[0]) <= 31 and ":" not in wb.sheetnames[0]

    def test_exportar_exige_login(self, client, db):
        resposta = client.get(
            reverse("reports:relatorio", args=["pesagens"]) + "?formato=xlsx"
        )
        assert resposta.status_code == 302


class TestInjecaoDeFormula:
    """Quem abre o arquivo no Excel executa `=...` como fórmula. Um comprador
    chamado `=HYPERLINK(...)` não pode virar ataque."""

    NOME_PERIGOSO = '=HYPERLINK("http://exemplo.invalido","clique")'

    @pytest.fixture
    def venda_para_comprador_perigoso(self, criar, confirmar, dados_abate):
        from apps.partners.models import Partner, PartnerRole, PartnerRoleChoice

        perigoso = Partner.objects.create(name=self.NOME_PERIGOSO)
        PartnerRole.objects.create(partner=perigoso, role=PartnerRoleChoice.FRIGORIFICO)
        return confirmar(criar(**{**dados_abate, "buyer": perigoso}))

    def test_csv_neutraliza_celula_de_texto_que_comeca_com_igual(
        self, client, gestor, season, venda_para_comprador_perigoso
    ):
        client.force_login(gestor)
        corpo = client.get(
            reverse("reports:relatorio", args=["vendas-e-abates"]) + "?formato=csv"
        ).content.decode()

        assert "'=HYPERLINK" in corpo
        # nenhuma célula começa com o "=" cru (depois de `;` ou início de linha)
        assert ";=HYPERLINK" not in corpo and "\r\n=HYPERLINK" not in corpo

    def test_xlsx_grava_como_texto_nunca_como_formula(
        self, client, gestor, season, venda_para_comprador_perigoso
    ):
        client.force_login(gestor)
        ws = abrir(
            client.get(
                reverse("reports:relatorio", args=["vendas-e-abates"]) + "?formato=xlsx"
            )
        ).active

        achadas = [
            c for lin in ws.iter_rows() for c in lin if "HYPERLINK" in str(c.value)
        ]
        assert achadas
        for celula in achadas:
            assert celula.data_type == "s"  # string, não "f" (fórmula)
            assert str(celula.value).startswith("'=")

    def test_numero_negativo_nao_e_confundido_com_formula(self):
        rel = services.Relatorio(
            titulo="x",
            descricao="",
            colunas=[
                services.Coluna("a", "Texto"),
                services.Coluna("b", "Valor", services.DINHEIRO),
            ],
            linhas=[{"a": "-teste", "b": D("-10.50")}],
        )
        corpo = services.csv_do_relatorio(rel)
        assert "'-teste" in corpo  # texto digitado: neutralizado
        assert "R$ -10,50" in corpo  # dinheiro: intocado
        assert "'R$" not in corpo


def test_a_previa_de_indicadores_exige_login(client, db):
    resposta = client.get(reverse("sales:previa"), {"head_count": "1"})
    assert resposta.status_code == 302
    assert resposta.url.startswith(reverse("accounts:login"))
