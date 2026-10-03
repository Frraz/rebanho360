"""Os formatos de saída: CSV, planilha (XLSX), JSON e PDF.

Todo escritor **grava em disco, linha a linha**: uma tabela de 200 mil linhas
não cabe na memória do processador junto com as outras. Cada conjunto passa
uma vez só pelo banco e alimenta todos os formatos pedidos ao mesmo tempo.

A planilha leva todos os conjuntos em abas de um arquivo só. CSV, JSON e PDF
levam um arquivo por conjunto. Relatórios prontos usam as funções que a tela
de relatórios já tem (mesmo número, mesmo formato) — ver `escrever_relatorio`.
"""

import csv
import json
import re
from dataclasses import dataclass
from datetime import date, datetime
from decimal import Decimal
from pathlib import Path

from django.template.loader import render_to_string
from django.utils import timezone

from apps.exports import tabular
from apps.exports.tabular import LEGIVEL, TECNICO

TIPOS = {
    "csv": "text/csv; charset=utf-8",
    "xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    "json": "application/json",
    "pdf": "application/pdf",
    "zip": "application/zip",
    "txt": "text/plain; charset=utf-8",
}

#: Versão do desenho dos arquivos. Sobe quando a estrutura do JSON, do
#: manifesto ou do PDF muda: quem guardou um backup sabe como lê-lo.
VERSAO_DO_FORMATO = "exportacao-v1"


@dataclass
class Arquivo:
    """Um arquivo gerado, ainda fora do pacote final."""

    caminho: Path
    nome: str  # caminho dentro do ZIP (ou nome do arquivo, se for único)
    tipo: str
    rotulo: str = ""
    registros: int | None = None
    aviso: str = ""


class Escritor:
    """Contrato: `abrir` → `linha`* → `fechar`, uma vez por conjunto; e
    `finalizar` uma vez no fim (a planilha só fecha aí)."""

    modo = LEGIVEL

    def abrir(self, conjunto, colunas) -> None:
        raise NotImplementedError

    def linha(self, valores) -> None:
        raise NotImplementedError

    def fechar(self) -> list[Arquivo]:
        return []

    def finalizar(self) -> list[Arquivo]:
        return []

    def abortar(self) -> None:
        """Fecha o que estiver aberto, sem pretender entregar nada."""


# --------------------------------------------------------------------------
# CSV
# --------------------------------------------------------------------------


class CsvEscritor(Escritor):
    """`br`: separador `;`, vírgula decimal, data dd/mm/aaaa e BOM (sem ele o
    Excel desfigura os acentos). Senão: `,`, ponto e ISO-8601, para programa."""

    modo = LEGIVEL

    def __init__(self, pasta: Path, *, br: bool):
        self.pasta, self.br = pasta, br
        self._arq = None

    def abrir(self, conjunto, colunas):
        self.conjunto, self.colunas, self.n = conjunto, colunas, 0
        self.caminho = self.pasta / "csv" / f"{conjunto.chave}.csv"
        self.caminho.parent.mkdir(parents=True, exist_ok=True)
        self._arq = open(
            self.caminho, "w", newline="", encoding="utf-8-sig" if self.br else "utf-8"
        )
        self._w = csv.writer(
            self._arq, delimiter=";" if self.br else ",", lineterminator="\r\n"
        )
        self._w.writerow([c.rotulo for c in colunas])

    def linha(self, valores):
        self._w.writerow(
            [
                tabular.para_csv(v, c.tipo, br=self.br)
                for v, c in zip(valores, self.colunas, strict=True)
            ]
        )
        self.n += 1

    def fechar(self):
        self.abortar()
        return [
            Arquivo(
                self.caminho,
                f"csv/{self.conjunto.chave}.csv",
                TIPOS["csv"],
                self.conjunto.titulo,
                self.n,
            )
        ]

    def abortar(self):
        if self._arq is not None:
            self._arq.close()
            self._arq = None


# --------------------------------------------------------------------------
# JSON
# --------------------------------------------------------------------------


def _coluna_para_dicionario(coluna, conjunto_do_modelo: dict) -> dict:
    descricao = {"chave": coluna.chave, "rotulo": coluna.rotulo, "tipo": coluna.tipo}
    if coluna.casas is not None:
        descricao["casas_decimais"] = coluna.casas
    if coluna.aponta_para:
        descricao["aponta_para"] = conjunto_do_modelo.get(
            coluna.aponta_para, coluna.aponta_para
        )
    return descricao


class JsonEscritor(Escritor):
    """Um objeto por conjunto, que se descreve sozinho (colunas e tipos) e traz
    os registros em `registros`. Valor técnico: vínculo é `id`, escolha é o
    código, `Decimal` é texto exato (nunca `float`)."""

    modo = TECNICO

    def __init__(self, pasta: Path, *, meta: dict, conjunto_do_modelo: dict):
        self.pasta, self.meta = pasta, meta
        self.conjunto_do_modelo = conjunto_do_modelo
        self._arq = None

    def abrir(self, conjunto, colunas):
        self.conjunto, self.colunas, self.n = conjunto, colunas, 0
        self.caminho = self.pasta / "json" / f"{conjunto.chave}.json"
        self.caminho.parent.mkdir(parents=True, exist_ok=True)
        self._arq = open(self.caminho, "w", encoding="utf-8")
        cabecalho = {
            "conjunto": conjunto.chave,
            "rotulo": conjunto.titulo,
            "formato": VERSAO_DO_FORMATO,
            "gerado_em": self.meta["gerado_em"],
            "filtros": self.meta["filtros"],
            "colunas": [
                _coluna_para_dicionario(c, self.conjunto_do_modelo) for c in colunas
            ],
        }
        texto = json.dumps(cabecalho, ensure_ascii=False, indent=1)
        self._arq.write(texto[:-2] + ',\n "registros": [\n')
        self._chaves = [c.chave for c in colunas]

    def linha(self, valores):
        registro = {
            chave: tabular.para_json(v) for chave, v in zip(self._chaves, valores)
        }
        self._arq.write(
            (",\n" if self.n else "")
            + json.dumps(registro, ensure_ascii=False, separators=(",", ":"))
        )
        self.n += 1

    def fechar(self):
        self._arq.write("\n ]\n}\n")
        self.abortar()
        return [
            Arquivo(
                self.caminho,
                f"json/{self.conjunto.chave}.json",
                TIPOS["json"],
                self.conjunto.titulo,
                self.n,
            )
        ]

    def abortar(self):
        if self._arq is not None:
            self._arq.close()
            self._arq = None


# --------------------------------------------------------------------------
# Planilha (XLSX)
# --------------------------------------------------------------------------

LIMITE_DA_CELULA = 32_000  # o Excel aceita 32.767 caracteres por célula
_NOME_DE_ABA_INVALIDO = re.compile(r"[\[\]:*?/\\]")


def _formato_de_numero(casas: int | None) -> str:
    return "0" if not casas else "0." + "0" * casas


class PlanilhaEscritor(Escritor):
    """Uma pasta de trabalho, uma aba por conjunto (mais a aba "Resumo").
    Os números são números (dá para somar e filtrar), as datas são datas."""

    modo = LEGIVEL

    def __init__(self, pasta: Path, *, meta: dict):
        import openpyxl

        self.pasta, self.meta = pasta, meta
        self.wb = openpyxl.Workbook(write_only=True)
        # A aba de resumo vem primeiro; é preenchida no fim, quando as
        # contagens já existem.
        self.resumo = self.wb.create_sheet("Resumo")
        self._abas: list[tuple[str, str, int]] = []
        self._nomes = {"Resumo"}
        self._ws = None

    def _nome_da_aba(self, titulo: str) -> str:
        base = _NOME_DE_ABA_INVALIDO.sub("", titulo)[:31] or "Dados"
        nome, i = base, 2
        while nome in self._nomes:
            sufixo = f" ({i})"
            nome, i = base[: 31 - len(sufixo)] + sufixo, i + 1
        self._nomes.add(nome)
        return nome

    def abrir(self, conjunto, colunas):
        from openpyxl.cell import WriteOnlyCell
        from openpyxl.styles import Alignment, Font
        from openpyxl.utils import get_column_letter

        self.conjunto, self.colunas, self.n = conjunto, colunas, 0
        self.nome = self._nome_da_aba(conjunto.titulo)
        self._ws = self.wb.create_sheet(self.nome)
        for i, coluna in enumerate(colunas, start=1):
            self._ws.column_dimensions[get_column_letter(i)].width = min(
                max(len(coluna.rotulo) + 2, 11), 40
            )
        self._ws.freeze_panes = "A2"
        cabecalho = []
        for coluna in colunas:
            celula = WriteOnlyCell(self._ws, value=coluna.rotulo)
            celula.font = Font(bold=True)
            celula.alignment = Alignment(
                horizontal="right" if coluna.eh_numerica() else "left", wrap_text=True
            )
            cabecalho.append(celula)
        self._ws.append(cabecalho)

    def _celula(self, valor, coluna):
        from openpyxl.cell import WriteOnlyCell
        from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE

        if valor is None:
            return None
        if isinstance(valor, bool):
            return "Sim" if valor else "Não"
        if isinstance(valor, Decimal):
            celula = WriteOnlyCell(self._ws, value=valor)
            celula.number_format = _formato_de_numero(coluna.casas)
            return celula
        if isinstance(valor, datetime):
            local = timezone.localtime(valor) if timezone.is_aware(valor) else valor
            celula = WriteOnlyCell(self._ws, value=local.replace(tzinfo=None))
            celula.number_format = "DD/MM/YYYY HH:MM"
            return celula
        if isinstance(valor, date):
            celula = WriteOnlyCell(self._ws, value=valor)
            celula.number_format = "DD/MM/YYYY"
            return celula
        if isinstance(valor, int):
            return valor
        if isinstance(valor, list | dict):
            valor = tabular._texto_estrutura(valor)
        texto = ILLEGAL_CHARACTERS_RE.sub("", str(valor))
        if len(texto) > LIMITE_DA_CELULA:
            texto = texto[: LIMITE_DA_CELULA - 40] + " …[texto cortado; veja o JSON]"
        # "=..." viraria fórmula: texto digitado por gente é neutralizado.
        return tabular.neutralizar_formula(texto)

    def linha(self, valores):
        self._ws.append(
            [self._celula(v, c) for v, c in zip(valores, self.colunas, strict=True)]
        )
        self.n += 1

    def fechar(self):
        self._abas.append((self.nome, self.conjunto.titulo, self.n))
        self._ws = None
        return []

    def finalizar(self):
        from openpyxl.cell import WriteOnlyCell
        from openpyxl.styles import Font

        if not self._abas:
            return []

        def negrito(texto):
            celula = WriteOnlyCell(self.resumo, value=texto)
            celula.font = Font(bold=True)
            return celula

        self.resumo.column_dimensions["A"].width = 34
        self.resumo.column_dimensions["B"].width = 44
        self.resumo.column_dimensions["C"].width = 14
        titulo = WriteOnlyCell(self.resumo, value="Rebanho360 — exportação de dados")
        titulo.font = Font(bold=True, size=14)
        self.resumo.append([titulo])
        self.resumo.append([f"Gerado em {self.meta['gerado_em_texto']}"])
        self.resumo.append([f"Por {self.meta['gerado_por']}"])
        self.resumo.append(["Filtros: " + " · ".join(self.meta["filtros"])])
        self.resumo.append([])
        self.resumo.append([negrito("Aba"), negrito("Conjunto"), negrito("Registros")])
        for aba, titulo_do_conjunto, n in self._abas:
            self.resumo.append([aba, titulo_do_conjunto, n])
        self.wb.properties.title = "Rebanho360 — exportação de dados"
        self.wb.properties.creator = "Rebanho360"
        caminho = self.pasta / "xlsx" / "dados.xlsx"
        caminho.parent.mkdir(parents=True, exist_ok=True)
        self.wb.save(caminho)
        return [
            Arquivo(
                caminho,
                "xlsx/dados.xlsx",
                TIPOS["xlsx"],
                "Planilha com uma aba por conjunto",
                sum(n for _, _, n in self._abas),
            )
        ]

    def abortar(self):
        try:
            self.wb.close()
        except Exception:  # pragma: no cover — só limpeza
            pass


# --------------------------------------------------------------------------
# PDF
# --------------------------------------------------------------------------


class _CabecalhoDoPdf:
    """O que o template de PDF dos relatórios espera de um `Relatorio`."""

    def __init__(self, titulo, descricao, notas, filtros):
        self.titulo, self.descricao = titulo, descricao
        self.notas, self.filtros = notas, filtros


class PdfEscritor(Escritor):
    """Um PDF por conjunto, com as **colunas principais** (o PDF é para ler,
    não para guardar: o resto está no CSV, na planilha e no JSON). Tabela
    grande vira vários PDFs de `LINHAS_POR_PARTE` linhas, até `LIMITE_DE_LINHAS`;
    o que passar disso fica de fora do PDF e o aviso diz quanto."""

    modo = LEGIVEL
    LINHAS_POR_PARTE = 2_000
    LIMITE_DE_LINHAS = 20_000

    def __init__(self, pasta: Path, *, meta: dict, ao_gerar_parte=None):
        self.pasta, self.meta = pasta, meta
        # Chamado a cada parte pronta: o PDF é a etapa lenta, e a tela precisa
        # saber que o processo continua vivo.
        self.ao_gerar_parte = ao_gerar_parte

    def abrir(self, conjunto, colunas):
        self.conjunto = conjunto
        self.todas = colunas
        self.escolhidas = tabular.colunas_do_pdf(conjunto, colunas)
        self.posicoes = [colunas.index(c) for c in self.escolhidas]
        self.buffer: list[list] = []
        self.partes: list[Path] = []
        self.recebidas = 0
        self.cortou = False
        self.total_de_colunas = len(colunas)

    def linha(self, valores):
        if self.recebidas >= self.LIMITE_DE_LINHAS:
            self.cortou = True
            return
        # Só descarrega quando chega a linha seguinte: assim a última parte
        # cheia não vira "parte 1 de 1" à toa.
        if len(self.buffer) >= self.LINHAS_POR_PARTE:
            self._descarregar()
        self.buffer.append(
            [tabular.para_pdf(valores[i], self.todas[i].tipo) for i in self.posicoes]
        )
        self.recebidas += 1

    def _descarregar(self, *, ultima=False):
        import weasyprint

        numero = len(self.partes) + 1
        unica = ultima and numero == 1
        primeira = (numero - 1) * self.LINHAS_POR_PARTE + 1
        notas = []
        if self.total_de_colunas > len(self.escolhidas):
            notas.append(
                f"Este PDF mostra as {len(self.escolhidas)} colunas principais de "
                f"{self.total_de_colunas}. O CSV, a planilha e o JSON trazem todas."
            )
        descricao = ""
        if not unica:
            descricao = f"Parte {numero} · linhas {primeira} a {primeira + len(self.buffer) - 1}"
        cabecalho = _CabecalhoDoPdf(
            self.conjunto.titulo, descricao, notas, self.meta["filtros"]
        )
        numericas = [c.eh_numerica() for c in self.escolhidas]
        tabela = {
            "cabecalho": [c.rotulo for c in self.escolhidas],
            "celulas": [list(zip(linha, numericas)) for linha in self.buffer],
            "indices_numericos": [i for i, n in enumerate(numericas) if n],
            "tem_totais": False,
        }
        html = render_to_string(
            "documents/relatorio.html",
            {
                "relatorio": cabecalho,
                "tabelas": [("", tabela)],
                "emitido_por": self.meta["gerado_por"],
                "emitido_em": self.meta["gerado_em_dt"],
                "versao": VERSAO_DO_FORMATO,
                "paisagem": len(self.escolhidas) > 7,
            },
        )
        sufixo = "" if unica else f"-parte-{numero:02d}"
        caminho = self.pasta / "pdf" / f"{self.conjunto.chave}{sufixo}.pdf"
        caminho.parent.mkdir(parents=True, exist_ok=True)
        weasyprint.HTML(string=html).write_pdf(target=str(caminho))
        self.partes.append(caminho)
        self.buffer = []
        if self.ao_gerar_parte:
            self.ao_gerar_parte()

    def fechar(self):
        if self.buffer or not self.partes:
            self._descarregar(ultima=True)
        aviso = ""
        if self.cortou:
            aviso = (
                f"O PDF de {self.conjunto.titulo} traz só as primeiras "
                f"{self.LIMITE_DE_LINHAS} linhas. O CSV, a planilha e o JSON têm todas."
            )
        arquivos = [
            Arquivo(
                caminho,
                f"pdf/{caminho.name}",
                TIPOS["pdf"],
                self.conjunto.titulo,
                None,
                aviso if i == 0 else "",
            )
            for i, caminho in enumerate(self.partes)
        ]
        return arquivos


# --------------------------------------------------------------------------
# Relatórios prontos
# --------------------------------------------------------------------------


def _relatorio_para_json(relatorio, meta: dict) -> str:
    def valor(v):
        return tabular.para_json(v)

    corpo = {
        "relatorio": relatorio.titulo,
        "descricao": relatorio.descricao,
        "formato": VERSAO_DO_FORMATO,
        "gerado_em": meta["gerado_em"],
        "filtros": list(relatorio.filtros),
        "notas": list(relatorio.notas),
        "colunas": [
            {"chave": c.chave, "rotulo": c.rotulo, "tipo": c.tipo}
            for c in relatorio.colunas
        ],
        "linhas": [
            {k: valor(v) for k, v in linha.items()} for linha in relatorio.linhas
        ],
    }
    if relatorio.totais is not None:
        corpo["totais"] = {k: valor(v) for k, v in relatorio.totais.items()}
    if relatorio.secoes:
        corpo["secoes"] = [
            json.loads(_relatorio_para_json(s, meta)) for s in relatorio.secoes
        ]
    return json.dumps(corpo, ensure_ascii=False, indent=1)


def escrever_relatorio(relatorio, slug: str, formatos, pasta: Path, meta: dict):
    """Os formatos pedidos de um relatório pronto, pelas mesmas funções que a
    tela de relatórios usa: o número do arquivo é o número da tela."""
    from apps.documents import services as documentos
    from apps.reports import services as relatorios

    arquivos = []
    destino = pasta / "relatorios"
    destino.mkdir(parents=True, exist_ok=True)

    def gravar(extensao, conteudo: bytes):
        caminho = destino / f"{slug}.{extensao}"
        caminho.write_bytes(conteudo)
        arquivos.append(
            Arquivo(
                caminho,
                f"relatorios/{slug}.{extensao}",
                TIPOS[extensao],
                relatorio.titulo,
            )
        )

    if "csv" in formatos:
        gravar("csv", relatorios.csv_do_relatorio(relatorio).encode("utf-8"))
    if "xlsx" in formatos:
        gravar(
            "xlsx",
            relatorios.xlsx_do_relatorio(
                relatorio,
                emitido_por=meta["gerado_por"],
                emitido_em=meta["gerado_em_dt"].replace(tzinfo=None),
            ),
        )
    if "json" in formatos:
        gravar("json", _relatorio_para_json(relatorio, meta).encode("utf-8"))
    if "pdf" in formatos:
        gravar(
            "pdf",
            documentos.renderizar_pdf(
                relatorio,
                emitido_por=meta["gerado_por"],
                emitido_em=meta["gerado_em_dt"],
                versao=VERSAO_DO_FORMATO,
            ),
        )
    return arquivos
